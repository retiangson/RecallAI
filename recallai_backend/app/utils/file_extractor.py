import io
import zipfile
import base64

import PyPDF2
import docx2txt
from pptx import Presentation
from openpyxl import load_workbook


def extract_text_from_any_file(file_bytes: bytes, filename: str) -> str:
    """
    Universal file extractor that works FROM RAW BYTES.
    Supports: PDF, DOCX, PPTX, XLSX, TXT, ZIP (recursively)
    """

    # Detect file extension
    ext = filename.lower().split(".")[-1]

    # -----------------------------
    # PDF
    # -----------------------------
    if ext == "pdf":
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text.strip()

    # -----------------------------
    # DOCX
    # -----------------------------
    if ext == "docx":
        return docx2txt.process(io.BytesIO(file_bytes)) or ""

    # -----------------------------
    # PPTX
    # -----------------------------
    if ext == "pptx":
        prs = Presentation(io.BytesIO(file_bytes))
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text.strip()

    # -----------------------------
    # XLSX
    # -----------------------------
    if ext == "xlsx":
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
        text = ""
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows():
                row_text = " | ".join([str(cell.value or "") for cell in row])
                text += row_text + "\n"
        return text.strip()

    # -----------------------------
    # ZIP (recursive extraction)
    # -----------------------------
    if ext == "zip":
        text = ""
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
            for name in z.namelist():
                if not z.getinfo(name).is_dir():
                    inner_bytes = z.read(name)
                    text += extract_text_from_any_file(inner_bytes, name) + "\n"
        return text.strip()

    # -----------------------------
    # Text files
    # -----------------------------
    if ext in ["txt", "md", "csv"]:
        return file_bytes.decode("utf-8", errors="ignore")

    # Unknown format → fallback: base64 for GPT
    return f"(Unrecognized file type)\n\nBase64:\n{base64.b64encode(file_bytes).decode()}"
